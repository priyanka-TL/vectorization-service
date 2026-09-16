import io
import uuid
import logging
from datetime import datetime
from typing import List, Dict, Any
from fastapi import HTTPException
from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document as LangchainDocument

from .base_processor import BaseFileProcessor
from app.config import settings
from app.services.translation_service import process_chunk

logger = logging.getLogger(__name__)

class DOCXProcessor(BaseFileProcessor):
    def get_supported_extensions(self) -> List[str]:
        return ['.docx', '.doc']

    async def process(self, file_content: bytes, filename: str, priority: str) -> List[Dict[str, Any]]:
        """Process DOCX file and return chunks"""
        self._validate_file_content(file_content, filename)

        try:
            # Read DOCX content
            doc = Document(io.BytesIO(file_content))
            text_content = ""

            # Extract text from paragraphs
            for para in doc.paragraphs:
                text_content += para.text + "\n"

            # Create document for LangChain
            langchain_doc = LangchainDocument(
                page_content=text_content,
                metadata={"source": filename, "type": "docx", "priority": priority},
            )

            # Split text into chunks
            text_splitter = RecursiveCharacterTextSplitter(
                chunk_size=settings.CHUNK_SIZE,
                chunk_overlap=settings.CHUNK_OVERLAP,
                length_function=len,
            )

            chunks = text_splitter.split_documents([langchain_doc])

            # Process chunks
            processed_chunks = []
            for chunk in chunks:
                chunk_id = uuid.uuid4().hex
                processed_text, is_hindi = process_chunk(
                    chunk.page_content, chunk_id
                )
                processed_chunk = {
                    "id": chunk_id,
                    "text": processed_text,
                    "metadata": {
                        **chunk.metadata,
                        "is_hindi": is_hindi,
                        "created_at": datetime.now().isoformat(),
                        "updated_at": datetime.now().isoformat(),
                    },
                }
                processed_chunks.append(processed_chunk)

            return processed_chunks

        except Exception as e:
            logger.error(f"DOCX processing error: {str(e)}")
            raise HTTPException(status_code=500, detail=f"Error processing DOCX: {str(e)}")
